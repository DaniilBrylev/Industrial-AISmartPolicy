"""
Генерация текста политики ИБ из анкеты и analysis_result (rules + AI).

CRUD сущностей PolicyDocument см. policy_document_service.
"""

from __future__ import annotations

import copy
import html
import json
import logging
import re
from pathlib import Path
from typing import Any

from app.core.config import Settings, settings as default_settings
from app.services.ai_service import AIService

logger = logging.getLogger(__name__)

_BULLET_LINE = re.compile(r"^\s*[-*•]\s+(.+)$")

# Ответы LLM, которые нельзя подставлять вместо фактических данных из analysis_result
_INSUFFICIENT_RE = re.compile(
    r"^\s*«?\s*недостаточно\s+данных\s*»?\s*\.?\s*$",
    re.IGNORECASE | re.MULTILINE,
)

# Удаление типовых LLM-заглушек в требованиях («не указан» и т.п.)
_REQ_PLACEHOLDER_RE = re.compile(
    r"\s*\([^)]*(?:не\s+указан|не\s+указано|n/?a)[^)]*\)",
    re.IGNORECASE,
)

_SEV_ORDER = {"low": 0, "medium": 1, "high": 2, "critical": 3}
_SEV_RU = {
    "low": "низкий",
    "medium": "средний",
    "high": "высокий",
    "critical": "критический",
}

# Типовые механизмы реализации угроз (детерминированный блок «модель угроз»).
THREAT_MECHANISM_BY_CODE: dict[str, str] = {
    "IT_DATA_LEAK": "несанкционированное копирование и вывод данных, эксплуатация каналов связи",
    "IT_UNAUTHORIZED_ACCESS": "компрометация учётных записей, ошибки конфигурации доступа, обход аутентификации",
    "IT_AVAILABILITY": "отказ или деградация сервисов, сбои инфраструктуры, дестабилизация доступности",
    "OT_PRODUCTION_STOP": "воздействие на управление технологическим процессом, останов производства",
    "OT_PROCESS_DISTURBANCE": "искажение технологических параметров и нарушение устойчивого режима",
    "OT_UNAUTHORIZED_ACCESS": "несанкционированное воздействие на средства АСУ ТП",
    "OT_SEGMENT_AVAILABILITY": "нарушение доступности сегментов OT-сети",
    "OT_CONTROLLER_COMPROMISE": "компрометация ПЛК/контроллеров и несанкционированное изменение логики",
    "OT_SETPOINT_TAMPERING": "несанкционированное изменение уставок и заданий",
    "OT_TELEMETRY_LOSS": "потеря или искажение телеметрии и данных мониторинга",
    "OT_CONTRACTOR_REMOTE": "злоупотребление удалённым доступом подрядчика в OT-контур",
    "OT_IT_OT_LATERAL": "распространение воздействия из IT в OT (lateral movement)",
    "OT_LEGACY_NO_PATCH": "эксплуатация уязвимостей при невозможности регулярного патчинга",
    "OT_SAFE_STATE_FAILURE": "невозможность или ошибка перевода процесса в безопасное состояние",
}

_PREVENT_MEASURE_HINT = re.compile(
    r"доступ|привилег|rbac|сегмент|изоляц|шифр|классификац|mfa|многофактор|"
    r"резерв|восстановлен|rto|rpo|патч|контроль\s+измен|компенс|dlp|dmz|межсет|vpn|разгранич|"
    r"инвентар|блокиров|физическ|jump|bastion|целостност|верси",
    re.IGNORECASE,
)

# Порядок и заголовки разделов экспорта (DOCX / HTML) — согласован с требованиями к структуре ПИБ.
POLICY_EXPORT_ORDER: list[tuple[str, str, str]] = [
    ("general", "1. Общие положения", "text"),
    ("scope", "2. Область применения", "text"),
    ("assets", "3. Активы", "list"),
    ("classification", "4. Классификация и среда", "list"),
    ("threat_model", "5. Модель угроз информационной безопасности", "list"),
    ("attacker_model", "6. Модель нарушителя информационной безопасности", "text"),
    ("risks", "7. Риски информационной безопасности", "list"),
    ("traceability_rows", "8. Связь «риск – требование – мера защиты»", "table"),
    ("requirements", "9. Требования по обеспечению информационной безопасности", "list"),
    ("prevention_measures", "10. Меры предотвращения угроз информационной безопасности", "list"),
    ("security_means", "11. Средства обеспечения информационной безопасности", "list"),
    ("measures", "12. Перечень мер защиты (компенсирующие и организационно-технические)", "list"),
    ("ib_system_description", "13. Описание системы обеспечения информационной безопасности", "text"),
    ("conclusion", "14. Заключение", "text"),
]


def _line_without_bullet_prefix(ln: str) -> str:
    s = ln.strip()
    m = _BULLET_LINE.match(s)
    if m:
        return m.group(1).strip()
    if re.match(r"^\d+[\.)]\s+", s):
        return re.sub(r"^\d+[\.)]\s+", "", s).strip()
    return s


def _is_insufficient_llm_response(text: str) -> bool:
    """True, если модель вернула только шаблонную отбивку без содержания."""
    if not (text or "").strip():
        return True
    lines = [ln for ln in text.strip().splitlines() if ln.strip()]
    if not lines:
        return True
    for ln in lines:
        core = _line_without_bullet_prefix(ln)
        if not core:
            continue
        if _INSUFFICIENT_RE.match(core) is None:
            return False
    return True


def _list_len(val: Any) -> int:
    return len(val) if isinstance(val, list) else 0


def _max_sev(a: str, b: str) -> str:
    return a if _SEV_ORDER.get(a, 0) >= _SEV_ORDER.get(b, 0) else b


def _sev_ru(sev: str) -> str:
    return _SEV_RU.get(sev, sev)


def _dedupe_lines_preserve_order(lines: list[str]) -> list[str]:
    seen: set[str] = set()
    out: list[str] = []
    for ln in lines:
        if ln not in seen:
            seen.add(ln)
            out.append(ln)
    return out


def _dedupe_risk_lines_from_analysis(normalized: dict[str, Any]) -> list[str]:
    risks = normalized.get("risks")
    if not isinstance(risks, list):
        return []
    by_code: dict[str, dict[str, Any]] = {}
    for item in risks:
        if not isinstance(item, dict):
            continue
        code = str(item.get("risk_code") or "").strip()
        if not code:
            continue
        aid = str(item.get("asset_id") or "").strip()
        title = str(item.get("title") or "").strip()
        sev = str(item.get("severity") or "medium").strip().lower()
        if sev not in _SEV_ORDER:
            sev = "medium"
        if code not in by_code:
            by_code[code] = {"title": title, "assets": set(), "sev": sev}
        if aid:
            by_code[code]["assets"].add(aid)
        if title and not by_code[code]["title"]:
            by_code[code]["title"] = title
        by_code[code]["sev"] = _max_sev(by_code[code]["sev"], sev)
    lines: list[str] = []
    for code in sorted(by_code.keys()):
        meta = by_code[code]
        assets = meta["assets"]
        ast_part = ", ".join(sorted(assets)) if assets else "—"
        title = meta["title"] or code
        lines.append(
            f"{code} — {title}; затрагиваемые активы: {ast_part}; "
            f"агрегированный уровень последствий: {_sev_ru(meta['sev'])}"
        )
    return lines


def _build_threat_model_lines(normalized: dict[str, Any]) -> list[str]:
    risks = normalized.get("risks")
    if not isinstance(risks, list):
        return []
    codes_order: list[str] = []
    seen: set[str] = set()
    titles: dict[str, str] = {}
    for item in risks:
        if not isinstance(item, dict):
            continue
        code = str(item.get("risk_code") or "").strip()
        if not code:
            continue
        if code not in titles or not titles[code]:
            titles[code] = str(item.get("title") or "").strip()
        if code not in seen:
            seen.add(code)
            codes_order.append(code)
    lines: list[str] = []
    for code in codes_order:
        mech = THREAT_MECHANISM_BY_CODE.get(
            code,
            "реализация угрозы возможна при наличии уязвимости или недостатка реализованных мер",
        )
        title = titles.get(code) or code
        lines.append(f"{code} — {title}. Типовой механизм: {mech}.")
    return lines


def _traceability_rows_from_links(links: Any) -> list[list[str]]:
    header = ["Идентификатор актива", "Риск", "Требование", "Мера защиты"]
    if not isinstance(links, list):
        return [header]
    seen: set[tuple[str, str, str, str]] = set()
    rows: list[list[str]] = [header]
    for ln in links:
        if not isinstance(ln, dict):
            continue
        row = (
            str(ln.get("asset_id") or "").strip(),
            str(ln.get("risk") or "").strip(),
            str(ln.get("requirement") or "").strip(),
            str(ln.get("measure") or "").strip(),
        )
        if not any(row) or row in seen:
            continue
        seen.add(row)
        rows.append(list(row))
    return rows


def _prevention_measures_subset(measures: list[str]) -> list[str]:
    prev = [m for m in measures if m and _PREVENT_MEASURE_HINT.search(m)]
    return prev if prev else list(measures)


def _security_means_catalog(measures: list[str]) -> list[str]:
    means: set[str] = set()
    joined = " ".join(measures).lower()
    if "dlp" in joined or "утечк" in joined:
        means.add("Средства предотвращения утечек информации (DLP)")
    if "siem" in joined or "корреляц" in joined:
        means.add("Средства SIEM и корреляции событий безопасности")
    if "журнал" in joined or "лог" in joined:
        means.add("Средства журналирования и централизованного хранения журналов")
    if "мониторинг" in joined or "оповещ" in joined:
        means.add("Средства мониторинга состояния и оповещения")
    if "mfa" in joined or "многофактор" in joined:
        means.add("Средства многофакторной аутентификации")
    if "шифр" in joined:
        means.add("Средства криптографической защиты информации")
    if "резерв" in joined or "rto" in joined or "rpo" in joined:
        means.add("Средства резервирования и восстановления")
    if "межсет" in joined or "dmz" in joined or "сегмент" in joined:
        means.add("Средства сетевой сегментации и межсетевого экранирования")
    if not means:
        means.add(
            "Конкретный перечень средств определяется локальным актом по результатам анализа рисков"
        )
    return sorted(means)


def _sanitize_requirement_line(text: str) -> str:
    t = _REQ_PLACEHOLDER_RE.sub("", str(text))
    t = re.sub(r"\s{2,}", " ", t).strip(" ;")
    return t


def _count_processes(response_data: dict[str, Any]) -> int:
    val = response_data.get("business_processes")
    return len(val) if isinstance(val, list) else 0


def _count_contractors(response_data: dict[str, Any]) -> int:
    for key in ("contractors", "external_contractors", "vendors"):
        val = response_data.get(key)
        if isinstance(val, list):
            return len(val)
    return 0


def _build_attacker_model_text(
    normalized: dict[str, Any], response_data: dict[str, Any]
) -> str:
    risks = normalized.get("risks")
    classified = normalized.get("classified_assets")
    risk_codes: set[str] = set()
    if isinstance(risks, list):
        for r in risks:
            if isinstance(r, dict):
                rc = str(r.get("risk_code") or "").strip()
                if rc:
                    risk_codes.add(rc)
    envs: set[str] = set()
    if isinstance(classified, list):
        for c in classified:
            if isinstance(c, dict):
                env = str(c.get("environment") or "").strip().upper()
                if env in ("IT", "OT"):
                    envs.add(env)

    contractors = _count_contractors(response_data)
    external_marker = any(
        x in risk_codes
        for x in ("IT_UNAUTHORIZED_ACCESS", "OT_UNAUTHORIZED_ACCESS", "OT_CONTRACTOR_REMOTE")
    )
    insider_marker = any(
        x in risk_codes for x in ("IT_DATA_LEAK", "IT_AVAILABILITY", "OT_SETPOINT_TAMPERING")
    )

    env_text = ", ".join(sorted(envs)) if envs else "IT"
    text_parts: list[str] = [
        (
            "Модель нарушителя формируется на основе состава активов, идентифицированных рисков и "
            f"операционного контекста подразделения. Для текущего профиля задействованы среды: {env_text}."
        ),
        (
            "Класс Н1 — внешний нарушитель низкой квалификации: массовые автоматизированные атаки, "
            "фишинг, эксплуатация известных уязвимостей."
        ),
    ]
    if external_marker:
        text_parts.append(
            "Класс Н2 — внешний целенаправленный нарушитель: компрометация удалённого доступа, "
            "попытки развития атаки и закрепления в инфраструктуре."
        )
    else:
        text_parts.append(
            "Класс Н2 — внешний целенаправленный нарушитель: целевые попытки обхода периметра "
            "и получения несанкционированного доступа."
        )
    if insider_marker:
        text_parts.append(
            "Класс Н3 — внутренний легитимный пользователь: ошибочные или злоупотребляющие действия "
            "в рамках предоставленных полномочий."
        )
    if "OT" in envs:
        text_parts.append(
            "Класс Н4 — внутренний привилегированный субъект (администратор/инженер АСУ ТП): "
            "несанкционированные изменения конфигураций и уставок, нарушение безопасного состояния."
        )
    else:
        text_parts.append(
            "Класс Н4 — внутренний привилегированный субъект (администратор): "
            "изменения конфигурации и журналирования в ущерб политике ИБ."
        )
    if contractors > 0:
        text_parts.append(
            f"В профиле анкеты зафиксировано внешнее взаимодействие (контрагентов: {contractors}), "
            "что усиливает значимость контроля удалённого доступа и договорных обязательств по ИБ."
        )
    text_parts.append(
        "Уточнение вероятностей реализации сценариев нарушителя выполняется при количественной оценке рисков."
    )
    return "\n\n".join(text_parts)


def _build_ib_system_description_text(
    normalized: dict[str, Any], response_data: dict[str, Any]
) -> str:
    assets = normalized.get("assets")
    classified = normalized.get("classified_assets")
    risks = normalized.get("risks")
    requirements = normalized.get("requirements")
    measures = normalized.get("measures")
    links = normalized.get("links")

    a_cnt = len(assets) if isinstance(assets, list) else 0
    c_cnt = len(classified) if isinstance(classified, list) else 0
    r_cnt = len(risks) if isinstance(risks, list) else 0
    req_cnt = len(requirements) if isinstance(requirements, list) else 0
    m_cnt = len(measures) if isinstance(measures, list) else 0
    l_cnt = len(links) if isinstance(links, list) else 0
    bp_cnt = _count_processes(response_data)
    contractors = _count_contractors(response_data)

    envs: set[str] = set()
    if isinstance(classified, list):
        for c in classified:
            if isinstance(c, dict):
                env = str(c.get("environment") or "").strip().upper()
                if env in ("IT", "OT"):
                    envs.add(env)
    env_text = ", ".join(sorted(envs)) if envs else "IT"

    scope_bits: list[str] = [
        f"среды эксплуатации: {env_text}",
    ]
    if a_cnt:
        scope_bits.append(
            f"инвентаризировано информационных активов: {a_cnt}"
        )
    if c_cnt:
        scope_bits.append(f"элементов классификации по средам и критичности: {c_cnt}")
    if bp_cnt:
        scope_bits.append(f"учтённых бизнес-процессов: {bp_cnt}")
    if r_cnt or req_cnt or m_cnt:
        scope_bits.append(
            f"профиль рисков и нормативных ожиданий включает "
            f"{r_cnt} запис(ей) о рисках, {req_cnt} требований и {m_cnt} мер защиты"
        )
    if l_cnt:
        scope_bits.append(
            f"настроена прослеживаемость «актив–риск–требование–мера» ({l_cnt} связей)"
        )
    scope_sentence = (
        "В контуре применения настоящей политики учитываются: "
        + "; ".join(scope_bits)
        + "."
    )

    para_intro = (
        "Система обеспечения информационной безопасности организации представляет собой совокупность "
        "взаимосвязанных организационных и технических мероприятий и средств, направленных на управление "
        "рисками ИБ и поддержание требуемого уровня защищённости информационных и технологических активов."
    )

    ot_extra = ""
    if "OT" in envs:
        ot_extra = (
            " Для объектов АСУ ТП и иных элементов OT-контура допускается применение компенсирующих мер "
            "при технологических ограничениях на обновление ПО и непрерывность технологического процесса."
        )

    contractor_extra = ""
    if contractors > 0:
        contractor_extra = (
            f" Учитывается взаимодействие с внешними контрагентами (в профиле: {contractors}); "
            "для таких сценариев усиливаются договорные требования ИБ, контроль удалённого доступа "
            "и разграничение ответственности."
        )

    para_subsystems = (
        "В составе системы обеспечения ИБ выделяются следующие функциональные подсистемы.\n\n"
        "1. Подсистема управления ИБ — установление политики и локальных регламентов, "
        "распределение ролей и ответственности, управление доступом и изменениями, планирование мер "
        "и контроль их исполнения."
        + contractor_extra
        + "\n\n"
        "2. Подсистема защиты — реализация превентивных и ограничивающих мер: модели доступа и "
        "разграничения полномочий, защита данных при хранении и передаче, сегментация и контроль "
        "сетевых взаимодействий, обеспечение целостности конфигураций."
        + ot_extra
        + "\n\n"
        "3. Подсистема мониторинга и обнаружения — централизованное журналирование, корреляция событий, "
        "выявление инцидентов и отклонений от установленных параметров безопасности и эксплуатации.\n\n"
        "4. Подсистема реагирования и восстановления — регламенты действий при инцидентах, перевод "
        "объектов в безопасное состояние (включая технологический контур при необходимости), "
        "резервирование критичных компонентов и восстановление в целевых значениях RTO/RPO."
    )

    para_cycle = (
        "Функционирование системы обеспечения ИБ организуется как непрерывный цикл: "
        "идентификация активов и контекста эксплуатации → анализ угроз и рисков → выбор и внедрение мер "
        "защиты → мониторинг, реагирование и учёт инцидентов → пересмотр политики и требований. "
        "Указанный цикл обеспечивает прослеживаемость между активами, рисками и мерами защиты и позволяет "
        "адаптировать систему ИБ к изменениям инфраструктуры и ландшафта угроз."
    )

    return "\n\n".join([para_intro, scope_sentence, para_subsystems, para_cycle])


def normalize_analysis_result(raw: Any) -> dict[str, Any]:
    """
    Приводит значение response_data['analysis_result'] к плоскому dict для генерации политики.

    Поддерживает:
    - JSON-строку вместо объекта;
    - вложенный объект report (как в QuestionnaireAnalyzeResponse: {\"report\": {...}});
    - опционально camelCase classifiedAssets → classified_assets.
    """
    if raw is None:
        return {}
    if isinstance(raw, str):
        s = raw.strip()
        if not s:
            return {}
        try:
            raw = json.loads(s)
        except json.JSONDecodeError:
            logger.warning("analysis_result: не удалось разобрать JSON-строку")
            return {}
    if not isinstance(raw, dict):
        return {}

    base = copy.deepcopy(raw)
    rep = base.get("report")
    if isinstance(rep, dict):
        for k in (
            "assets",
            "classified_assets",
            "risks",
            "requirements",
            "measures",
            "links",
            "warnings",
            "ai_enrichment",
            "traceability_map",
            "analysis_meta",
        ):
            if k in rep:
                base[k] = rep[k]

    if "classified_assets" not in base and isinstance(base.get("classifiedAssets"), list):
        base["classified_assets"] = base["classifiedAssets"]

    logger.info(
        "normalize_analysis_result: len(assets)=%s len(classified_assets)=%s len(risks)=%s "
        "len(requirements)=%s len(measures)=%s",
        _list_len(base.get("assets")),
        _list_len(base.get("classified_assets")),
        _list_len(base.get("risks")),
        _list_len(base.get("requirements")),
        _list_len(base.get("measures")),
    )
    return base


def transform_analysis_to_policy_sections(
    normalized_analysis: dict[str, Any],
    policy_svc: PolicyService | None = None,
) -> dict[str, list[str]]:
    """
    Плоские списки строк для разделов политики (вход — результат normalize_analysis_result).

    Маппинг: assets, classified_assets → classification, risks, requirements, measures.
    """
    svc = policy_svc or PolicyService()
    return {
        "assets": svc._lines_from_report_list(normalized_analysis.get("assets"), "asset"),
        "classification": svc._lines_classified(normalized_analysis.get("classified_assets")),
        "risks": svc._lines_from_report_list(normalized_analysis.get("risks"), "risk"),
        "requirements": svc._string_list(normalized_analysis.get("requirements")),
        "measures": svc._string_list(normalized_analysis.get("measures")),
    }


class PolicyService:
    """Сборка структуры политики, вызов AI по разделам, HTML и DOCX."""

    def __init__(self, app_settings: Settings | None = None) -> None:
        self._settings = app_settings or default_settings
        raw = getattr(self._settings, "policy_export_dir", "var/generated_policies")
        self._export_dir = Path(raw)

    def docx_filename(self, questionnaire_id: int) -> str:
        return f"policy_q{questionnaire_id}.docx"

    def docx_path(self, questionnaire_id: int) -> Path:
        return self._export_dir / self.docx_filename(questionnaire_id)

    def build_policy_structure(
        self, response_data: dict[str, Any], analysis_result: dict[str, Any]
    ) -> dict[str, Any]:
        """
        Собирает черновую структуру до вызова LLM.
        Списки — краткие строки; general/scope/conclusion — текстовые наброски.
        """
        dp = response_data.get("department_profile")
        if not isinstance(dp, dict):
            dp = {}

        notes = response_data.get("additional_notes")
        notes_str = notes if isinstance(notes, str) else ""

        general = (
            f"Подразделение: описание — {dp.get('description', '')}; "
            f"руководитель — {dp.get('manager_name', '')}; "
            f"контакты — {dp.get('contact_info', '')}."
        ).strip()

        scope = (
            "Область действия политики формируется на основе учтённых активов, "
            "бизнес-процессов и заявленного контекста организации. "
            f"Дополнительные сведения: {notes_str or 'не указаны'}."
        )

        analysis_norm = normalize_analysis_result(analysis_result)
        sections = transform_analysis_to_policy_sections(analysis_norm, self)
        assets_lines = list(sections["assets"])
        class_lines = _dedupe_lines_preserve_order(list(sections["classification"]))
        risk_lines = _dedupe_risk_lines_from_analysis(analysis_norm)
        req_lines = list(sections["requirements"])
        meas_lines = list(sections["measures"])

        if not assets_lines and isinstance(response_data.get("assets"), list):
            assets_lines = self._lines_from_raw_assets(response_data["assets"])

        threat_lines = _build_threat_model_lines(analysis_norm)
        trace_rows = _traceability_rows_from_links(analysis_norm.get("links"))
        prevention_lines = _prevention_measures_subset(meas_lines)
        security_means_lines = _security_means_catalog(meas_lines)

        conclusion = (
            "Политика носит рамочный характер; конкретные регламенты и сроки внедрения "
            "определяются локальными актами на основе результатов анализа рисков."
        )

        def _or_missing(lines: list[str]) -> list[str]:
            return lines if lines else ["данные отсутствуют"]

        attacker_model_text = _build_attacker_model_text(analysis_norm, response_data)
        ib_system_description_text = _build_ib_system_description_text(
            analysis_norm, response_data
        )

        return {
            "general": general,
            "scope": scope,
            "assets": _or_missing(assets_lines),
            "classification": _or_missing(class_lines),
            "threat_model": _or_missing(threat_lines),
            "attacker_model": attacker_model_text,
            "risks": _or_missing(risk_lines),
            "traceability_rows": trace_rows,
            "requirements": _or_missing(req_lines),
            "prevention_measures": _or_missing(prevention_lines),
            "security_means": _or_missing(security_means_lines),
            "measures": _or_missing(meas_lines),
            "ib_system_description": ib_system_description_text,
            "conclusion": conclusion,
        }

    def _string_list(self, val: Any) -> list[str]:
        if isinstance(val, list):
            return [str(x).strip() for x in val if str(x).strip()]
        return []

    def _lines_from_report_list(self, val: Any, kind: str) -> list[str]:
        if not isinstance(val, list):
            return []
        out: list[str] = []
        for i, item in enumerate(val):
            if isinstance(item, dict):
                if kind == "asset":
                    out.append(
                        f"{item.get('id', i)}: {item.get('name', '')} "
                        f"({item.get('asset_type', '')}), критичность {item.get('criticality', '')}"
                    )
                elif kind == "risk":
                    out.append(
                        f"{item.get('risk_code', '')} — {item.get('title', '')} "
                        f"(актив {item.get('asset_id', '')}, severity {item.get('severity', '')})"
                    )
                else:
                    out.append(json.dumps(item, ensure_ascii=False))
            else:
                out.append(str(item))
        return out

    def _lines_classified(self, val: Any) -> list[str]:
        if not isinstance(val, list):
            return []
        lines: list[str] = []
        for item in val:
            if not isinstance(item, dict):
                lines.append(str(item))
                continue
            aid = item.get("asset_id", "")
            lines.append(
                f"{aid}: среда {item.get('environment', '')}, "
                f"критичность исх. {item.get('original_criticality', '')} → "
                f"эфф. {item.get('effective_criticality', '')}, "
                f"процессов: {item.get('usage_count', 0)}"
            )
        return lines

    def _lines_from_raw_assets(self, assets: list[Any]) -> list[str]:
        lines: list[str] = []
        for a in assets:
            if not isinstance(a, dict):
                continue
            lines.append(
                f"{a.get('id')}: {a.get('name', '')} ({a.get('asset_type', '')}) "
                f"— {a.get('criticality', '')}"
            )
        return lines

    async def generate_policy_text(
        self, structure: dict[str, Any], ai_service: AIService
    ) -> dict[str, Any]:
        """
        Для каждого логического раздела вызывает AIService.generate_policy_section.
        Детерминированно (без LLM): модель угроз, таблица трассировки, меры предотвращения, перечень средств ИБ.
        Гибридно (черновик из правил + уточнение LLM): модель нарушителя, описание системы обеспечения ИБ.
        """
        out: dict[str, Any] = {}

        deterministic_lists = (
            "threat_model",
            "prevention_measures",
            "security_means",
        )
        hybrid_text_keys = ("attacker_model", "ib_system_description")
        for key in hybrid_text_keys:
            out[key] = str(structure.get(key) or "")

        tr = structure.get("traceability_rows")
        out["traceability_rows"] = (
            copy.deepcopy(tr) if isinstance(tr, list) else [["", "", "", ""]]
        )

        for key in deterministic_lists:
            items = structure.get(key)
            out[key] = list(items) if isinstance(items, list) else []

        text_keys = ("general", "scope", "conclusion", "attacker_model", "ib_system_description")
        list_keys = ("assets", "classification", "risks", "requirements", "measures")
        ai_context_base = {k: structure[k] for k in ("general", "scope", "conclusion") if k in structure}

        for key in text_keys:
            draft = str(structure.get(key, "") or "")
            payload = {
                "section": key,
                "draft": draft,
                "full_context": structure,
            }
            try:
                generated = await ai_service.generate_policy_section(payload)
                g = (generated or "").strip()
                if _is_insufficient_llm_response(g):
                    out[key] = draft
                else:
                    out[key] = g or draft
            except Exception as e:  # noqa: BLE001
                logger.error("AI section %s failed, using draft: %s", key, e)
                out[key] = draft

        for key in list_keys:
            items = structure.get(key)
            if not isinstance(items, list):
                items = []
            payload = {
                "section": key,
                "items": items,
                "full_context": ai_context_base,
            }
            try:
                generated = await ai_service.generate_policy_section(payload)
                out[key] = self._text_to_bullet_list(generated, fallback=items)
            except Exception as e:  # noqa: BLE001
                logger.error("AI section %s failed, using structure items: %s", key, e)
                out[key] = list(items)

        if isinstance(out.get("requirements"), list):
            out["requirements"] = [
                _sanitize_requirement_line(str(x)) for x in out["requirements"]
            ]

        return out

    def _text_to_bullet_list(self, text: str, *, fallback: list[str]) -> list[str]:
        if not (text or "").strip() or _is_insufficient_llm_response(text):
            return list(fallback)
        lines = []
        for line in text.splitlines():
            line = line.strip()
            if not line:
                continue
            m = _BULLET_LINE.match(line)
            if m:
                lines.append(m.group(1).strip())
            elif re.match(r"^\d+[\.)]\s+", line):
                lines.append(re.sub(r"^\d+[\.)]\s+", "", line).strip())
            else:
                lines.append(line)
        return lines if lines else list(fallback)

    def render_html_preview(self, policy: dict[str, Any]) -> str:
        """Простой HTML для предпросмотра (экранирование текста)."""
        parts: list[str] = [
            "<!DOCTYPE html><html><head><meta charset=\"utf-8\"><title>Политика ИБ</title>"
            "<style>table{border-collapse:collapse;width:100%;margin:0.5em 0}"
            "th,td{border:1px solid #444;padding:6px;text-align:left}</style></head><body>"
        ]
        parts.append("<h1>Проект политики информационной безопасности</h1>")

        for key, title, kind in POLICY_EXPORT_ORDER:
            parts.append(f"<h2>{html.escape(title)}</h2>")
            val = policy.get(key)
            if kind == "table":
                rows = val if isinstance(val, list) else []
                if len(rows) < 2:
                    parts.append(
                        "<p><em>Детализированные связи не сформированы (нет данных анализа).</em></p>"
                    )
                else:
                    parts.append("<table><tbody>")
                    for i, row in enumerate(rows):
                        if not isinstance(row, list):
                            continue
                        parts.append("<tr>")
                        tag = "th" if i == 0 else "td"
                        for cell in row:
                            parts.append(f"<{tag}>{html.escape(str(cell))}</{tag}>")
                        parts.append("</tr>")
                    parts.append("</tbody></table>")
            elif kind == "list":
                if not isinstance(val, list):
                    val = []
                parts.append("<ul>")
                for item in val:
                    parts.append(f"<li>{html.escape(str(item))}</li>")
                parts.append("</ul>")
            else:
                text = str(val or "")
                for para in text.split("\n\n"):
                    p = para.strip()
                    if p:
                        parts.append(f"<p>{html.escape(p)}</p>")

        parts.append("</body></html>")
        return "\n".join(parts)

    def export_docx(self, policy: dict[str, Any], file_path: str) -> None:
        try:
            from docx import Document
        except ImportError as e:  # pragma: no cover
            raise RuntimeError("Установите пакет python-docx") from e

        path = Path(file_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()
        doc.add_heading("Политика информационной безопасности", level=0)

        for key, heading, kind in POLICY_EXPORT_ORDER:
            doc.add_heading(heading, level=1)
            val = policy.get(key)
            if kind == "table":
                rows = val if isinstance(val, list) else []
                if len(rows) < 2:
                    doc.add_paragraph(
                        "Детализированные связи не сформированы (нет данных анализа)."
                    )
                else:
                    ncols = max((len(r) for r in rows if isinstance(r, list)), default=4)
                    tbl = doc.add_table(rows=len(rows), cols=ncols)
                    tbl.style = "Table Grid"
                    for i, row in enumerate(rows):
                        if not isinstance(row, list):
                            continue
                        for j in range(ncols):
                            cell_text = str(row[j]) if j < len(row) else ""
                            tbl.rows[i].cells[j].text = cell_text
            elif kind == "list":
                items = val if isinstance(val, list) else []
                for item in items:
                    p = doc.add_paragraph(style="List Bullet")
                    p.add_run(str(item))
            else:
                text = str(val or "")
                for para in text.split("\n\n"):
                    p = para.strip()
                    if p:
                        doc.add_paragraph(p)

        doc.save(str(path))
        logger.info("DOCX exported: %s", path)


def get_policy_service() -> PolicyService:
    return PolicyService()
